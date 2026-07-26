#!/usr/bin/env python3
"""Build a branded reference.docx for pandoc: AMIU navy/gold institutional style."""
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

SRC = "/home/user/shroyalschools/assets/reference-base.docx"
OUT = "/home/user/shroyalschools/assets/amiu-reference.docx"

NAVY = RGBColor(0x12, 0x2A, 0x4E)
NAVY_DK = RGBColor(0x0A, 0x18, 0x30)
GOLD = RGBColor(0xB0, 0x86, 0x25)
TEXT = RGBColor(0x1A, 0x1F, 0x2B)
GREY = RGBColor(0x5B, 0x63, 0x72)

# Three-role editorial system, replacing the earlier four-family setup
# (Cinzel/Cormorant Garamond/EB Garamond/Liberation Sans) that a close read
# correctly flagged as visually flat: three of those four were the same
# old-style-Garamond-descended serif at different weights, and the small-
# caps/kicker/caption/table voice throughout was Liberation Sans — the
# open-source Arial substitute, i.e. the single most "this is a Word
# document" signal on every page. Each role below is now a genuinely
# distinct family, not just a different size of the same one:
#   Fraunces        — display: covers, monumental numerals, headings, pull
#                      quotes. An ink-trap display serif with real range
#                      (Black through Light Italic) and a contemporary
#                      editorial identity, not the "wedding invitation"
#                      register Cinzel/Cormorant read as.
#   Source Serif 4   — reading: body prose. Built for long-form reading,
#                      deliberately NOT another Garamond derivative, so it
#                      reads as a distinct second voice against Fraunces
#                      rather than a same-family cousin.
#   Archivo          — structural: captions, sub-subsection headings,
#                      table data, TOC rows, numbered-provision labels.
#                      A confident contemporary grotesk (Black through
#                      Regular) replacing Liberation Sans everywhere.
#
# Typography-audit correction: ceremonial kickers/eyebrows (the "ARTICLE
# IV · SECTION 4.1" tags above every heading, the "SECTIONS OF THIS
# ARTICLE" labels on dividers, and the running header/footer) were
# originally tracked all-caps Archivo — the same device countless SaaS
# pitch decks and nonprofit annual reports use for wayfinding labels, and
# the one place a genuine "this reads corporate, not Oxford/Cambridge"
# critique landed on an otherwise-sound system. Real university-press and
# FT-style editorial typography marks this register with small caps OF
# THE SERIF FAMILY, not a second grotesk voice — so MONUMENTAL_FONT now
# points at Source Serif 4 (small caps, not all caps) instead of Archivo.
HEAD_FONT = "Source Serif 4"
BODY_FONT = "Archivo"
SERIF_FONT = "Liberation Serif"
MONUMENTAL_FONT = "Source Serif 4"         # Level 1 — ceremonial kickers/running heads (serif small caps,
                                            # not a grotesk sans — see the typography-audit note below)
DISPLAY_FONT = "Fraunces"                  # Level 2 — chapter titles, founding statements, pull quotes
DISPLAY_SEMIBOLD = "Fraunces SemiBold"     # Level 3 — section headings
DISPLAY_BLACK = "Fraunces Black"           # monumental numerals and cover hero type

doc = Document(SRC)

def get_style(document, name):
    """Robust lookup bypassing python-docx's lowercase 'heading N' alias bug."""
    for s in document.styles:
        if s.name == name:
            return s
    raise KeyError(name)

def has_style(document, name):
    return any(s.name == name for s in document.styles)

def set_font(style, name=BODY_FONT, size=11, color=TEXT, bold=False, italic=False,
             all_caps=False, small_caps=False, spacing=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    # Pandoc's default reference template points Heading/Title styles at
    # theme fonts (asciiTheme="majorHAnsi", etc). When both a theme
    # reference and an explicit font coexist on the same rFonts element,
    # LibreOffice's headless docx->pdf conversion resolves the theme font
    # instead of honoring the explicit override — so the theme attributes
    # must be removed, not just shadowed, whenever an explicit font is set.
    for theme_attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        qattr = qn(theme_attr)
        if rFonts.get(qattr) is not None:
            del rFonts.attrib[qattr]
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:cs'), name)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.font.italic = italic
    style.font.all_caps = all_caps
    style.font.small_caps = small_caps
    # Explicit language tag so LibreOffice's hyphenation engine resolves the
    # correct dictionary (hyph_en_US.dic) rather than silently skipping
    # hyphenation for lack of a language association.
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), 'en-US')
    if spacing is not None:
        rpr2 = style.element.get_or_add_rPr()
        spc = OxmlElement('w:spacing')
        spc.set(qn('w:val'), str(spacing))
        rpr2.append(spc)

def set_para_fmt(style, space_before=0, space_after=8, line=1.15, keep_next=False,
                  align=None, border_bottom=None, border_color=None, widow_control=None,
                  no_hyphens=False):
    pf = style.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align
    if widow_control is not None:
        pf.widow_control = widow_control
    if no_hyphens:
        # Headings/display lines should never hyphenate — a mid-word break
        # in a title reads as an error, not elegance. Document-wide
        # autoHyphenation is on for body prose; this opts individual
        # styles back out.
        pPr = style.element.get_or_add_pPr()
        supp = OxmlElement('w:suppressAutoHyphens')
        supp.set(qn('w:val'), 'true')
        pPr.append(supp)
    if border_bottom:
        pPr = style.element.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), str(border_bottom))
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), border_color or '122A4E')
        pBdr.append(bottom)
        pPr.append(pBdr)

# ---------------- Normal / body ----------------
# Source Serif 4 at 10.8pt / ~13.8pt leading (ratio 1.28). Source Serif 4's
# x-height reads noticeably larger than EB Garamond's at the same nominal
# size, so the point size drops (11.5 -> 10.8) to hold the same apparent
# body-text scale rather than let the font swap alone inflate the page count.
normal = get_style(doc, 'Normal')
set_font(normal, HEAD_FONT, 10.8, TEXT)
set_para_fmt(normal, 0, 9, 1.28, widow_control=True)

# ---------------- Title (used for the doc's H1-level "Title" para if any) ----------------
title = get_style(doc, 'Title')
set_font(title, DISPLAY_FONT, 34, NAVY, bold=True, spacing=6)
set_para_fmt(title, 0, 6, 1.05, align=WD_ALIGN_PARAGRAPH.LEFT, no_hyphens=True)

if has_style(doc, 'Subtitle'):
    subtitle = get_style(doc, 'Subtitle')
    set_font(subtitle, DISPLAY_FONT, 17, GOLD, italic=True)
    set_para_fmt(subtitle, 0, 20, 1.15, no_hyphens=True)

# ---------------- Heading 1 ----------------
# NOTE: page breaks before H1 headings are handled by EXPLICIT raw-openxml
# breaks in the markdown, not by a style-level pageBreakBefore. H1 is reused
# across Part dividers *and* front/back-matter section openers with very
# different page-break needs, and stacking a style-level break under an
# already-explicit break produced stray blank pages during QA.
h1 = get_style(doc, 'Heading 1')
set_font(h1, DISPLAY_FONT, 32, NAVY, bold=True, spacing=4)
set_para_fmt(h1, 4, 16, 1.05, keep_next=True, border_bottom=18, border_color='B08625', no_hyphens=True)

# ---------------- Heading 2 (Sections) ----------------
# NOTE: space-before is intentionally small — the new SectionKicker paragraph
# immediately above every Section heading now owns the larger "before" gap,
# so the kicker and its heading read as one tightly-coupled unit.
h2 = get_style(doc, 'Heading 2')
set_font(h2, DISPLAY_SEMIBOLD, 22, NAVY, bold=True, spacing=2)
set_para_fmt(h2, 2, 8, 1.08, keep_next=True, border_bottom=6, border_color='D8DCE5', no_hyphens=True)

# ---------------- Custom style: ceremonial message byline ----------------
# Elevates the "From the Chairman, Board of Trustees" / "By the Supreme
# Strategic Planning Council" speaker lines on the Foreword and the three
# Message pages above ordinary italic body text — paired with the drop
# cap, this gives those four ceremonial pages a distinct opening register
# instead of running straight from the H1 title into body-sized prose.
if not has_style(doc, 'MessageByline'):
    byline = doc.styles.add_style('MessageByline', WD_STYLE_TYPE.PARAGRAPH)
    byline.base_style = get_style(doc, 'Normal')
else:
    byline = get_style(doc, 'MessageByline')
set_font(byline, DISPLAY_FONT, 15, GOLD, italic=True, spacing=4)
set_para_fmt(byline, 6, 18, 1.2, no_hyphens=True)

# ---------------- Custom style: Section kicker badge (Part · Section N) ----------------
if not has_style(doc, 'SectionKicker'):
    sec_kicker = doc.styles.add_style('SectionKicker', WD_STYLE_TYPE.PARAGRAPH)
    sec_kicker.base_style = get_style(doc, 'Normal')
else:
    sec_kicker = get_style(doc, 'SectionKicker')
set_font(sec_kicker, MONUMENTAL_FONT, 10.5, GOLD, bold=True, italic=True, small_caps=True, spacing=18)
set_para_fmt(sec_kicker, 22, 2, 1.1, keep_next=True, no_hyphens=True)

# ---------------- Custom style: Part-divider kicker label ----------------
if not has_style(doc, 'PartKicker'):
    kicker = doc.styles.add_style('PartKicker', WD_STYLE_TYPE.PARAGRAPH)
    kicker.base_style = get_style(doc, 'Normal')
else:
    kicker = get_style(doc, 'PartKicker')
set_font(kicker, MONUMENTAL_FONT, 12, GOLD, bold=True, italic=True, small_caps=True, spacing=24)
set_para_fmt(kicker, 0, 2, 1.1, no_hyphens=True)

# ---------------- Custom style: Part-divider thesis line ----------------
if not has_style(doc, 'PartThesis'):
    thesis = doc.styles.add_style('PartThesis', WD_STYLE_TYPE.PARAGRAPH)
    thesis.base_style = get_style(doc, 'Normal')
else:
    thesis = get_style(doc, 'PartThesis')
set_font(thesis, DISPLAY_FONT, 16.5, NAVY, italic=True)
set_para_fmt(thesis, 4, 18, 1.35, no_hyphens=True)

# ---------------- Heading 3 ----------------
h3 = get_style(doc, 'Heading 3')
set_font(h3, BODY_FONT, 12, GOLD, bold=True, all_caps=False)
set_para_fmt(h3, 14, 6, 1.1, keep_next=True, no_hyphens=True)

# ---------------- Heading 4 ----------------
if has_style(doc, 'Heading 4'):
    h4 = get_style(doc, 'Heading 4')
    set_font(h4, BODY_FONT, 10.5, NAVY, bold=True, italic=True)
    set_para_fmt(h4, 10, 4, 1.1, keep_next=True, no_hyphens=True)

# ---------------- TOC / Contents heading ----------------
for nm in ['TOC Heading']:
    if has_style(doc, nm):
        tocH = get_style(doc, nm)
        set_font(tocH, DISPLAY_FONT, 26, NAVY, bold=True, spacing=3)
        set_para_fmt(tocH, 0, 14, 1.1, no_hyphens=True)

for lvl, sz in zip(['TOC 1', 'TOC 2', 'TOC 3'], [12, 10.8, 10]):
    if has_style(doc, lvl):
        st = get_style(doc, lvl)
        set_font(st, BODY_FONT, sz, TEXT, bold=(lvl == 'TOC 1'))
        set_para_fmt(st, 4 if lvl == 'TOC 1' else 1, 4 if lvl == 'TOC 1' else 1, 1.1)

# ---------------- Block quote (pull quotes) ----------------
if has_style(doc, 'Block Text'):
    bq = get_style(doc, 'Block Text')
    set_font(bq, DISPLAY_FONT, 16, NAVY, italic=True)
    set_para_fmt(bq, 10, 10, 1.3)
if has_style(doc, 'Quote'):
    bq = get_style(doc, 'Quote')
    set_font(bq, DISPLAY_FONT, 16, NAVY, italic=True)
    set_para_fmt(bq, 10, 10, 1.3)

# ---------------- Caption ----------------
if has_style(doc, 'Caption'):
    cap = get_style(doc, 'Caption')
    set_font(cap, BODY_FONT, 10, GOLD, bold=True, italic=False, small_caps=True, spacing=6)
    set_para_fmt(cap, 6, 16, 1.15, no_hyphens=True)

# ---------------- Table styles ----------------
if has_style(doc, 'Table Grid'):
    tg = get_style(doc, 'Table Grid')
    try:
        set_font(tg, BODY_FONT, 9.5, TEXT)
    except Exception:
        pass

# ---------------- Compact / First paragraph ----------------
for nm in ['Compact', 'Body Text', 'First Paragraph']:
    if has_style(doc, nm):
        st = get_style(doc, nm)
        set_font(st, HEAD_FONT, 10.8, TEXT)
        set_para_fmt(st, 0, 7, 1.25, widow_control=True)

# ==================================================================
# Page setup: US Letter, generous but efficient margins, section props
# ==================================================================
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.left_margin = Inches(1.1)
section.right_margin = Inches(1.1)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.header_distance = Inches(0.5)
section.footer_distance = Inches(0.5)
section.different_first_page_header_footer = True

# ==================================================================
# Header (title running header) & Footer (page number + confidentiality)
# ==================================================================
def add_field(paragraph, field_code, font=BODY_FONT, size=8.5, color=GREY, bold=False):
    run = paragraph.add_run()
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = field_code
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    return run

# --- Default header (all pages after first) ---
header = section.header
header.is_linked_to_previous = False
hp = header.paragraphs[0]
hp.text = ""
hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT)
hp.paragraph_format.space_after = Pt(0)
r1 = hp.add_run("Al-Mulk International University")
r1.font.name = MONUMENTAL_FONT; r1.font.size = Pt(9); r1.font.color.rgb = NAVY; r1.font.bold = False
r1.font.italic = True
r1.font.small_caps = True
rpr1 = r1._r.get_or_add_rPr()
spc1 = OxmlElement('w:spacing'); spc1.set(qn('w:val'), '8'); rpr1.append(spc1)
r2 = hp.add_run("\t")
r3 = hp.add_run("Strategic Implementation Blueprint 2028–2050")
r3.font.name = DISPLAY_FONT; r3.font.size = Pt(9.5); r3.font.color.rgb = GREY; r3.font.italic = True
# rule under header — a true hairline in restrained gold (0.375pt): thin
# enough not to compete with the content, and a genuine hairline reads as
# restraint rather than the "gold overuse" a thicker rule on every page
# would be.
pPr = hp._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '3')
bottom.set(qn('w:space'), '6'); bottom.set(qn('w:color'), 'B08625')
pBdr.append(bottom)
pPr.append(pBdr)

# --- Default footer (all pages after first) ---
# Minimalist three-part layout: volume identifier at left, nothing at
# center, page number alone at right (no "of NNN" — a running total reads
# as a progress bar, not a page number). The confidentiality/status
# notice lives once, on the Copyright & Publication Notice page.
footer = section.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
fp.text = ""
fp.paragraph_format.tab_stops.add_tab_stop(Inches(3.15), WD_TAB_ALIGNMENT.CENTER)
fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT)
fp.paragraph_format.space_before = Pt(6)
r4 = fp.add_run("AMIU-SB-002")
r4.font.name = MONUMENTAL_FONT; r4.font.size = Pt(7.5); r4.font.color.rgb = GREY
r4.font.small_caps = True
rpr4 = r4._r.get_or_add_rPr()
spc4 = OxmlElement('w:spacing'); spc4.set(qn('w:val'), '10'); rpr4.append(spc4)
fp.add_run("\t\t")
add_field(fp, 'PAGE', font=HEAD_FONT, size=10, color=NAVY, bold=False)

# --- First-page header/footer: blank (cover page has no running head) ---
fph = section.first_page_header
fph.is_linked_to_previous = False
fph.paragraphs[0].text = ""
fpf = section.first_page_footer
fpf.is_linked_to_previous = False
fpf.paragraphs[0].text = ""

# ==================================================================
# Executive table style: navy header row, gold rule, subtle zebra banding
# ==================================================================
def qn_(tag):
    return qn(tag)

table_style = None
for s in doc.styles.element.findall(qn('w:style')):
    if s.get(qn('w:styleId')) == 'Table':
        table_style = s
        break

if table_style is not None:
    tblPr = table_style.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        table_style.append(tblPr)
    # Clean horizontal-rule borders (no verticals) at table level
    borders = OxmlElement('w:tblBorders')
    for edge, sz, color in [('top', 8, '122A4E'), ('bottom', 8, '122A4E'),
                             ('insideH', 4, 'D8DCE5')]:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    for edge in ['left', 'right', 'insideV']:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tblPr.append(borders)

    # Remove existing firstRow tblStylePr (will be replaced) and any band defs
    for tsp in table_style.findall(qn('w:tblStylePr')):
        table_style.remove(tsp)

    def make_tblStylePr(kind, fill=None, text_color=None, bold=False, border_bottom=None):
        tsp = OxmlElement('w:tblStylePr')
        tsp.set(qn('w:type'), kind)
        if bold or text_color:
            rPr = OxmlElement('w:rPr')
            if bold:
                rPr.append(OxmlElement('w:b'))
            if text_color:
                c = OxmlElement('w:color'); c.set(qn('w:val'), text_color); rPr.append(c)
            tsp.append(rPr)
        tcPr = OxmlElement('w:tcPr')
        if fill:
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
            tcPr.append(shd)
        if border_bottom:
            tcBorders = OxmlElement('w:tcBorders')
            b = OxmlElement('w:bottom')
            b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '12'); b.set(qn('w:color'), border_bottom)
            tcBorders.append(b)
            tcPr.append(tcBorders)
        vAlign = OxmlElement('w:vAlign'); vAlign.set(qn('w:val'), 'center')
        tcPr.append(vAlign)
        tcMar = OxmlElement('w:tcMar')
        for side, val in [('top', 80), ('bottom', 80), ('left', 120), ('right', 120)]:
            m = OxmlElement(f'w:{side}')
            m.set(qn('w:w'), str(val)); m.set(qn('w:type'), 'dxa')
            tcMar.append(m)
        tcPr.append(tcMar)
        tsp.append(tcPr)
        return tsp

    # NOTE: table-style rPr color is beaten by the explicit color set on the
    # "Compact" paragraph style pandoc applies to every cell, so header
    # legibility is achieved via a light tint + bold rather than reversed text.
    table_style.append(make_tblStylePr('firstRow', fill='DCE3F0',
                                        bold=True, border_bottom='122A4E'))
    table_style.append(make_tblStylePr('band1Horz', fill='F6F7FA'))

doc.save(OUT)
print("Saved", OUT)
